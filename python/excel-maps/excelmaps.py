from openpyxl import load_workbook
import geocoder
import osrm
from tkinter import Tk
from tkinter.filedialog import askopenfilename, asksaveasfilename
from tkinter.messagebox import showinfo
from functools import lru_cache

START_COL = 'K'
DEST_COL = 'F'
DIST_COL = 'N'
TIME_COL = 'O'
SHEET = 'MATRICE'

@lru_cache(maxsize=None)
def get_coords(address):
    return geocoder.osm(address).latlng

@lru_cache(maxsize=None)
def get_dist_time(coords1, coords2):
    res = osrm.simple_route(coords1, coords2)
    print(res)
    return None, None

def get_files():
    Tk().withdraw()
    in_file = askopenfilename(title='Open Excel File', defaultextension='.xlsx', filetypes=(('Excel','.xlsx .xls'),))
    out_file = asksaveasfilename(title='Save Excel File', defaultextension='.xlsx', filetypes=(('Excel','.xlsx .xls'),))
    return in_file, out_file

def main():
    in_file, out_file = get_files()
    wb = load_workbook(filename = in_file)
    ws = wb.get_sheet_by_name(SHEET)

    start_col = map(lambda v: v.value, ws[START_COL][1:])
    dest_col = map(lambda v: v.value, ws[DEST_COL][1:])
    dist_col = ws[DIST_COL][1:]
    time_col = ws[TIME_COL][1:]

    done = 0
    skipped = 0
    for start, dest, dist, time in zip(start_col, dest_col, dist_col, time_col):
        if start and dest and not dist.value and not time.value:
            dist.value, time.value = get_dist_time(get_coords(start), get_coords(dest))
            done += 1
        else:
            skipped += 1

    wb.save(filename = out_file)
    showinfo(title='Completed', message=f'Computed {done}/{done+skipped}, Excel file saved')

if __name__ == '__main__':
    main()
